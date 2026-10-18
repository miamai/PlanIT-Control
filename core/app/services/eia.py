from __future__ import annotations

import csv
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import httpx
from openpyxl import load_workbook
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.ext.asyncio import AsyncSession
from timezonefinder import TimezoneFinder

from app.models.entities import Site
from app.providers.base import RetryingHTTPProvider

EIA_RELEASE = "August 2026"
EIA_SOURCE_URL = "https://www.eia.gov/electricity/data/eia860m/xls/august_generator2026.xlsx"
EIA_API_URL = "https://api.eia.gov/v2/electricity/operating-generator-capacity/data"
EIA_API_METADATA_URL = "https://api.eia.gov/v2/electricity/operating-generator-capacity"
CACHE_PATH = Path(__file__).resolve().parents[2] / "data" / "eia860m_august_2026_sites.csv"


@dataclass
class PlantRecord:
    eia_plant_id: int
    name: str
    latitude: float
    longitude: float
    state: str
    county: str | None
    operator_name: str | None
    technologies: dict[str, float] = field(default_factory=lambda: defaultdict(float))
    energy_sources: dict[str, float] = field(default_factory=lambda: defaultdict(float))

    @property
    def plant_type(self) -> str | None:
        return max(self.technologies, key=self.technologies.get) if self.technologies else None

    @property
    def primary_energy_source(self) -> str | None:
        return (
            max(self.energy_sources, key=self.energy_sources.get) if self.energy_sources else None
        )


def _valid_coordinates(latitude: object, longitude: object) -> bool:
    try:
        lat, lon = float(latitude), float(longitude)
        return -90 <= lat <= 90 and -180 <= lon <= 180 and not (lat == 0 and lon == 0)
    except (TypeError, ValueError):
        return False


def parse_eia_workbook(path: Path) -> list[PlantRecord]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Operating" not in workbook.sheetnames:
        raise ValueError("EIA workbook does not contain an Operating sheet")
    sheet = workbook["Operating"]
    rows = sheet.iter_rows(values_only=True)
    next(rows, None)
    next(rows, None)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    required = {
        "Entity Name",
        "Plant ID",
        "Plant Name",
        "Plant State",
        "County",
        "Technology",
        "Energy Source Code",
        "Nameplate Capacity (MW)",
        "Latitude",
        "Longitude",
    }
    missing = required.difference(headers)
    if missing:
        raise ValueError(f"EIA workbook is missing columns: {sorted(missing)}")
    index = {name: headers.index(name) for name in required}
    plants: dict[int, PlantRecord] = {}
    for row in rows:
        if not _valid_coordinates(row[index["Latitude"]], row[index["Longitude"]]):
            continue
        try:
            plant_id = int(row[index["Plant ID"]])
        except (TypeError, ValueError):
            continue
        capacity_raw = row[index["Nameplate Capacity (MW)"]]
        try:
            capacity = float(capacity_raw or 0)
        except (TypeError, ValueError):
            capacity = 0
        record = plants.setdefault(
            plant_id,
            PlantRecord(
                eia_plant_id=plant_id,
                name=str(row[index["Plant Name"]]).strip(),
                latitude=float(row[index["Latitude"]]),
                longitude=float(row[index["Longitude"]]),
                state=str(row[index["Plant State"]]).strip(),
                county=(str(row[index["County"]]).strip() or None),
                operator_name=(str(row[index["Entity Name"]]).strip() or None),
            ),
        )
        technology = str(row[index["Technology"]] or "").strip()
        energy_source = str(row[index["Energy Source Code"]] or "").strip()
        if technology:
            record.technologies[technology] += capacity
        if energy_source:
            record.energy_sources[energy_source] += capacity
    return sorted(plants.values(), key=lambda plant: plant.eia_plant_id)


def parse_cached_sites(path: Path = CACHE_PATH) -> list[PlantRecord]:
    records: list[PlantRecord] = []
    with path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            record = PlantRecord(
                eia_plant_id=int(row["eia_plant_id"]),
                name=row["name"],
                latitude=float(row["latitude"]),
                longitude=float(row["longitude"]),
                state=row["state"],
                county=row["county"] or None,
                operator_name=row["operator_name"] or None,
            )
            if row["plant_type"]:
                record.technologies[row["plant_type"]] = 1
            if row["primary_energy_source"]:
                record.energy_sources[row["primary_energy_source"]] = 1
            records.append(record)
    return records


def select_reference_sites(
    records: list[PlantRecord], *, per_state: int = 3, maximum: int = 153
) -> list[PlantRecord]:
    """Choose a bounded, deterministic reference set from operating plants."""
    by_state: dict[str, list[PlantRecord]] = defaultdict(list)
    for record in records:
        by_state[record.state].append(record)
    selected: dict[int, PlantRecord] = {}
    for state in sorted(by_state):
        ranked = sorted(
            by_state[state],
            key=lambda item: (-sum(item.technologies.values()), item.eia_plant_id),
        )
        for record in ranked[:per_state]:
            selected[record.eia_plant_id] = record
    for record in sorted(records, key=lambda item: item.eia_plant_id):
        if len(selected) >= maximum:
            break
        selected.setdefault(record.eia_plant_id, record)
    return [selected[key] for key in sorted(selected)][:maximum]


async def fetch_eia_api_records(
    api_key: str,
    *,
    timeout_seconds: float = 60.0,
    page_size: int = 5000,
    transport: httpx.AsyncBaseTransport | None = None,
) -> list[PlantRecord]:
    """Load generator rows from EIA APIv2 and aggregate them by plant."""
    async with httpx.AsyncClient(timeout=timeout_seconds, transport=transport) as client:
        http = RetryingHTTPProvider(client)
        metadata = await http.get_json(EIA_API_METADATA_URL, params={"api_key": api_key})
        route = metadata.get("response", metadata)
        fields = {row.get("id") for row in route.get("data", []) if isinstance(row, dict)}
        required = {
            "plant_id",
            "plant_name",
            "stateid",
            "county",
            "latitude",
            "longitude",
            "entity_name",
            "technology",
            "energy_source_code",
            "nameplate_capacity_mw",
            "status",
        }
        if not required.issubset(fields):
            raise ValueError(f"EIA API metadata is missing fields: {sorted(required - fields)}")
        params = {
            "api_key": api_key,
            "data[]": sorted(required - {"status"}),
            "facets[status][]": "OP",
            "length": page_size,
            "offset": 0,
        }
        plants: dict[int, PlantRecord] = {}
        while True:
            payload = await http.get_json(EIA_API_URL, params=params)
            response = payload.get("response", payload)
            rows = response.get("data", [])
            for row in rows:
                if str(row.get("status", "OP")).upper() not in {"OP", "OPERATING"}:
                    continue
                if not _valid_coordinates(row.get("latitude"), row.get("longitude")):
                    continue
                try:
                    plant_id = int(row["plant_id"])
                    capacity = float(row.get("nameplate_capacity_mw") or 0)
                    latitude = float(row["latitude"])
                    longitude = float(row["longitude"])
                except (KeyError, TypeError, ValueError):
                    continue
                record = plants.setdefault(
                    plant_id,
                    PlantRecord(
                        eia_plant_id=plant_id,
                        name=str(row.get("plant_name") or "").strip(),
                        latitude=latitude,
                        longitude=longitude,
                        state=str(row.get("stateid") or "").strip(),
                        county=str(row.get("county") or "").strip() or None,
                        operator_name=str(row.get("entity_name") or "").strip() or None,
                    ),
                )
                technology = str(row.get("technology") or "").strip()
                source = str(row.get("energy_source_code") or "").strip()
                if technology:
                    record.technologies[technology] += capacity
                if source:
                    record.energy_sources[source] += capacity
            total = int(response.get("total", len(rows)))
            if not rows or int(params["offset"]) + len(rows) >= total:
                break
            params["offset"] = int(params["offset"]) + len(rows)
        return sorted(plants.values(), key=lambda item: item.eia_plant_id)


async def download_eia_workbook(destination: Path) -> Path:
    async with httpx.AsyncClient(timeout=60, follow_redirects=True) as client:
        response = await client.get(EIA_SOURCE_URL)
        response.raise_for_status()
        destination.write_bytes(response.content)
    return destination


async def import_plants(
    session: AsyncSession,
    records: list[PlantRecord],
    *,
    source_url: str = EIA_SOURCE_URL,
    source_release: str = EIA_RELEASE,
) -> int:
    timezone_finder = TimezoneFinder(in_memory=True)
    count = 0
    for record in records:
        timezone = timezone_finder.timezone_at(lat=record.latitude, lng=record.longitude) or "UTC"
        values = {
            "eia_plant_id": record.eia_plant_id,
            "name": record.name,
            "latitude": record.latitude,
            "longitude": record.longitude,
            "state": record.state,
            "county": record.county,
            "plant_type": record.plant_type,
            "primary_energy_source": record.primary_energy_source,
            "operator_name": record.operator_name,
            "timezone": timezone,
            "source_release": source_release,
            "source_url": source_url,
            "source_metadata": {
                "dataset": "EIA-860M",
                "all_technologies": sorted(record.technologies),
                "all_energy_sources": sorted(record.energy_sources),
                "reference_location_notice": "Not a verified active construction site",
            },
            "is_reference_location": True,
        }
        statement = insert(Site).values(**values)
        statement = statement.on_conflict_do_update(
            index_elements=["eia_plant_id"],
            set_={key: value for key, value in values.items() if key != "eia_plant_id"},
        )
        await session.execute(statement)
        count += 1
    await session.commit()
    return count
